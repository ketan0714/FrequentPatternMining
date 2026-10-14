from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import shutil, zipfile, os

BASE=Path('/mnt/data/Apriori_DHP_New')
GRAPHS=BASE/'graphs'
OUT=BASE/'Apriori_vs_DHP_Attractive.pptx'
NAVY=RGBColor(16,36,58); NAVY2=RGBColor(24,51,76); TEAL=RGBColor(18,184,166); CORAL=RGBColor(240,123,98); GOLD=RGBColor(230,182,85); BLUE=RGBColor(106,143,183); WHITE=RGBColor(255,255,255); LIGHT=RGBColor(242,246,249); MID=RGBColor(101,119,136); DARK=RGBColor(31,49,66); PALE=RGBColor(226,237,241)
prs=Presentation(); prs.slide_width= Inches(13.333); prs.slide_height= Inches(7.5)
blank=prs.slide_layouts[6]

def bg(slide, color=WHITE):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=color

def shape(slide,x,y,w,h,fill, radius=True, line=None):
    typ=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s=slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    s.line.fill.background() if line is None else None
    if line:
        s.line.color.rgb=line
    if radius:
        try: s.adjustments[0]=0.12
        except: pass
    return s

def txt(slide,x,y,w,h,text,size=18,color=DARK,bold=False,font='Aptos',align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def rich_lines(slide,x,y,w,h,items,size=16,color=DARK,spacing=8):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.05); tf.margin_right=Inches(.05); tf.margin_top=Inches(.05); tf.margin_bottom=Inches(.03)
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item; p.level=0; p.font.name='Aptos'; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(spacing); p.bullet=True
    return box

def header(slide,title,kicker=None,dark=False):
    c=WHITE if dark else NAVY
    if kicker: txt(slide,.62,.35,4,.25,kicker.upper(),10,TEAL if dark else BLUE,True)
    txt(slide,.62,.68,12.1,.55,title,27,c,True)
    shape(slide,.62,1.32,1.05,.055,TEAL,False)

def footer(slide,n,dark=False):
    c=RGBColor(183,199,211) if dark else MID
    txt(slide,.62,7.14,8,.18,'APRIORI  /  DHP  •  TECHNICAL COMPARISON',8,c,True)
    txt(slide,12.1,7.1,.6,.22,f'{n:02d}',9,c,True,align=PP_ALIGN.RIGHT)

def card(slide,x,y,w,h,title,body,accent=TEAL,dark=False):
    shape(slide,x,y,w,h,NAVY2 if dark else LIGHT,True)
    shape(slide,x,y,.075,h,accent,False)
    txt(slide,x+.22,y+.18,w-.4,.34,title,16,WHITE if dark else NAVY,True)
    txt(slide,x+.22,y+.62,w-.42,h-.76,body,12,RGBColor(218,229,236) if dark else MID)

def add_img(slide,path,x,y,w,h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

def newslide(title,kicker=None,dark=False):
    s=prs.slides.add_slide(blank); bg(s,NAVY if dark else WHITE); header(s,title,kicker,dark); footer(s,len(prs.slides),dark); return s

# 1 title
s=prs.slides.add_slide(blank); bg(s,NAVY)
shape(s,8.55,0,4.78,7.5,NAVY2,False)
for i,(x,y,w,c) in enumerate([(9.0,1.0,1.1,TEAL),(10.45,1.65,.7,CORAL),(11.35,.8,.45,GOLD),(9.65,3.0,.65,BLUE),(11.0,3.65,1.0,TEAL),(9.0,5.1,.5,GOLD),(11.6,5.4,.65,CORAL)]):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(w)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background()
for x,y,w,h in [(9.25,2.35,2.7,.07),(9.4,4.5,2.6,.07),(10.0,1.0,.07,3.8),(11.6,2.2,.07,3.2)]: shape(s,x,y,w,h,RGBColor(66,94,116),False)
txt(s,.8,.75,6.8,.35,'DATA MINING  /  ALGORITHM BENCHMARK',11,TEAL,True)
txt(s,.8,1.55,7.3,1.45,'Apriori vs DHP',38,WHITE,True)
txt(s,.8,3.02,7.0,.85,'Hash-based pruning and transaction trimming for frequent itemset mining',21,RGBColor(209,224,233))
shape(s,.8,4.35,1.2,.07,CORAL,False)
txt(s,.8,4.75,6.7,.65,'Implementation • Synthetic benchmark • Performance trade-offs',15,WHITE,True)
txt(s,.8,6.55,6.8,.35,'Updated to match the supplied Python dataset and experiment code',10,RGBColor(183,199,211))
footer(s,1,True)

# 2 overview
s=newslide('What this comparison asks','01  /  ORIENTATION')
card(s,.7,1.7,3.85,2.0,'Baseline: Apriori','Generate candidate itemsets level by level; prune candidates whose subsets are not frequent.',TEAL)
card(s,4.75,1.7,3.85,2.0,'Modification: DHP','Use pair-hash bucket counts to safely prune some 2-itemset candidates before exact counting.',CORAL)
card(s,8.8,1.7,3.85,2.0,'Additional variant','Apply a simplified transaction-trimming step to reduce later-pass data work.',BLUE)
txt(s,.85,4.35,11.5,.45,'Evaluation dimensions',19,NAVY,True)
for i,(t,b,c) in enumerate([('Runtime','Elapsed time by dataset size and support',TEAL),('Memory','Peak traced allocation after data is loaded',CORAL),('Correctness','Same frequent itemsets/counts across variants',BLUE)]):
    x=.8+i*4.15; shape(s,x,4.95,3.75,1.15,LIGHT); txt(s,x+.2,5.12,3.35,.25,t,14,c,True); txt(s,x+.2,5.48,3.35,.42,b,11,MID)

# 3 basics
s=newslide('Association-rule mining in one slide','02  /  CORE CONCEPTS')
card(s,.75,1.65,3.75,1.75,'Transaction','A basket or event containing a set of items, e.g. {bread, milk, eggs}.',TEAL)
card(s,4.8,1.65,3.75,1.75,'Frequent itemset','An item combination whose support reaches the chosen minimum threshold.',CORAL)
card(s,8.85,1.65,3.75,1.75,'Association rule','A directional implication X → Y evaluated with support and confidence.',BLUE)
txt(s,.9,3.9,5.2,.4,'Key measures',18,NAVY,True)
shape(s,.9,4.45,5.25,1.5,LIGHT); txt(s,1.15,4.7,4.8,.35,'Support(X) = transactions containing X / all transactions',14,NAVY,True); txt(s,1.15,5.22,4.8,.35,'Confidence(X → Y) = support(X ∪ Y) / support(X)',13,MID)
txt(s,6.8,3.9,5.4,.4,'Where it is used',18,NAVY,True)
rich_lines(s,6.85,4.42,5.4,1.8,['Retail basket analysis and cross-selling','Web clickstream and content recommendations','Diagnosis/event co-occurrence and inventory planning'],14,DARK,10)
txt(s,.95,6.55,11.6,.3,'Scope note: this implementation benchmarks frequent itemset mining; rule generation/confidence reporting is described conceptually, not implemented in the benchmark script.',10,CORAL,True)

# 4 Apriori principle
s=newslide('Apriori: the downward-closure principle','03  /  BASELINE')
shape(s,.9,1.7,5.25,3.85,LIGHT)
txt(s,1.2,1.98,4.7,.4,'If an itemset is frequent…',17,NAVY,True)
for x,y,label,c in [(1.5,2.75,'{A,B,C}',TEAL),(1.3,3.75,'{A,B}',BLUE),(2.7,3.75,'{A,C}',BLUE),(4.1,3.75,'{B,C}',BLUE)]:
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.25), Inches(.58)); sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background(); txt(s,x,y+.12,1.25,.25,label,13,WHITE,True,align=PP_ALIGN.CENTER)
for x1,y1,x2,y2 in [(2.1,3.33,1.9,3.72),(2.1,3.33,3.3,3.72),(2.1,3.33,4.7,3.72)]:
    line=s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); line.line.color.rgb=MID; line.line.width=Pt(1.6)
txt(s,6.8,1.8,5.3,.5,'Candidate generation = join + prune',19,NAVY,True)
rich_lines(s,6.85,2.45,5.25,2.5,['Join frequent (k−1)-itemsets with a shared prefix.','Prune a k-itemset if any required (k−1)-subset is absent.','Count support only for the remaining candidates.'],15,DARK,12)
shape(s,6.85,5.25,5.25,.78,PALE); txt(s,7.08,5.45,4.8,.3,'Main cost: candidate explosion + repeated database scans',13,NAVY,True)

# 5 Apriori flow
s=newslide('Apriori workflow','03  /  BASELINE')
steps=[('1','Count singletons'),('2','Keep frequent L₁'),('3','Join + prune to C₂'),('4','Scan & count support'),('5','Repeat until no candidates')]
for i,(num,label) in enumerate(steps):
    x=.75+i*2.5
    shape(s,x,2.3,2.1,1.35,LIGHT,True)
    shape(s,x+.15,2.48,.42,.42,TEAL if i<2 else BLUE,True)
    txt(s,x+.15,2.55,.42,.2,num,12,WHITE,True,align=PP_ALIGN.CENTER)
    txt(s,x+.18,3.05,1.75,.45,label,12,NAVY,True,align=PP_ALIGN.CENTER)
    if i<4:
        line=s.shapes.add_connector(1, Inches(x+2.12), Inches(2.98), Inches(x+2.42), Inches(2.98)); line.line.color.rgb=MID; line.line.width=Pt(2)
shape(s,1.0,4.55,11.2,1.0,NAVY,True)
txt(s,1.3,4.83,10.6,.4,'Output: all frequent itemsets and their exact support counts',18,WHITE,True,align=PP_ALIGN.CENTER)
txt(s,1.05,6.0,11.1,.35,'Apriori remains the reference implementation used to compare output equality.',12,MID,False,align=PP_ALIGN.CENTER)

# 6 DHP concept
s=newslide('DHP: use hash buckets to prune pairs early','04  /  MODIFICATION')
card(s,.75,1.7,3.65,2.0,'Pass 1','Count each item and hash every observed pair into a bucket table.',TEAL)
card(s,4.85,1.7,3.65,2.0,'Safe pruning rule','If a bucket count is below min support, every pair mapped there is infrequent.',CORAL)
card(s,8.95,1.7,3.65,2.0,'Exact verification','A bucket above threshold is only a possibility; candidate pairs still need exact support counting.',BLUE)
shape(s,1.0,4.35,11.2,1.25,LIGHT,True)
txt(s,1.3,4.58,10.6,.35,'Hash used in the supplied code',16,NAVY,True,align=PP_ALIGN.CENTER)
txt(s,1.3,5.05,10.6,.35,'bucket(a,b) = (a × 1009 + b) mod B     •     default B = 20,011',16,CORAL,True,align=PP_ALIGN.CENTER)
txt(s,1.05,6.1,11.0,.3,'Hash collisions can keep extra candidates, but they cannot make a truly frequent pair disappear.',11,MID,False,align=PP_ALIGN.CENTER)

# 7 variants
s=newslide('Three implementations compared','04  /  IMPLEMENTATION')
card(s,.8,1.7,3.8,2.15,'Apriori','Shared candidate join/prune and exact support counting. No hash pruning or transaction trimming.',TEAL)
card(s,4.8,1.7,3.8,2.15,'DHP — hash only','Adds first-pass pair-bucket counts; uses the same candidate generator and exact support counter.',CORAL)
card(s,8.8,1.7,3.8,2.15,'DHP — hash + trim','Adds a simplified transaction-trimming step after passes to reduce later-pass records.',BLUE)
shape(s,.9,4.55,11.4,1.05,PALE,True)
txt(s,1.2,4.82,10.8,.4,'Important implementation detail: this is a clear educational Python version, not a reproduction of every optimization in the original DHP paper.',13,NAVY,True,align=PP_ALIGN.CENTER)
txt(s,1.05,6.0,11.1,.35,'The code asserts output equality against Apriori for the completed configurations and runs a brute-force self-test at startup.',11,MID)

# 8 dataset/method
s=newslide('Dataset and benchmark design','05  /  EXPERIMENT')
# compact table
rows=[('Item universe','200 item IDs'),('Observed distinct items','190 / 194 / 198 across 2k / 5k / 10k'),('Average transaction length','≈ 11.3 items'),('Dataset origin','Synthetic Quest-style generator; fixed seed 2026'),('Size sweep','2,000 • 5,000 • 10,000 transactions'),('Support sweep','5,000 transactions; 2% and 1% support'),('Bucket sweep','5,000 transactions; B = 997, 4,999, 20,011')]
for i,(a,b) in enumerate(rows):
    y=1.65+i*.62
    shape(s,.8,y,3.05,.48,NAVY2,True); txt(s,1.0,y+.11,2.7,.22,a,11,WHITE,True)
    shape(s,4.0,y,8.4,.48,LIGHT,True); txt(s,4.2,y+.11,7.95,.22,b,11,DARK)
txt(s,.9,6.25,11.5,.35,'Timing shown in this deck is a single completed run per configuration; treat small differences as indicative, not statistically conclusive.',10,CORAL,True)

# 9 time graph
s=newslide('Result 1 — runtime grows with dataset size','06  /  RESULTS')
add_img(s,GRAPHS/'graph1_time_vs_size.png',.55,1.55,8.5,4.8)
card(s,9.15,1.85,3.35,1.35,'10k transactions','Apriori: 0.804 s\nDHP hash only: 0.919 s\nDHP + trim: 1.034 s',TEAL)
card(s,9.15,3.45,3.35,1.35,'Interpretation','On this synthetic workload, hash overhead and trimming cost outweigh pruning savings in elapsed time.',CORAL)
txt(s,9.25,5.25,3.1,.8,'All three returned 4,370 frequent itemsets at 10k / 1% support.',11,MID)

# 10 memory graph
s=newslide('Result 2 — memory trade-off','06  /  RESULTS')
add_img(s,GRAPHS/'graph2_memory_vs_size.png',.55,1.55,8.5,4.8)
card(s,9.15,1.85,3.35,1.35,'Hash only','Lower peak traced memory than Apriori in these measured cases.',CORAL)
card(s,9.15,3.45,3.35,1.35,'Hash + trim','Memory is higher here; the extra bucket table and trimmed transaction structures have a cost.',BLUE)
txt(s,9.25,5.25,3.1,.85,'Memory uses tracemalloc peak after input loading—not total process RSS.',10,MID)

# 11 support graph
s=newslide('Result 3 — lower support means more work','06  /  RESULTS')
add_img(s,GRAPHS/'graph3_time_vs_support.png',.55,1.55,8.5,4.8)
card(s,9.15,1.85,3.35,1.35,'At 2% support','1,188 frequent itemsets; fewer candidates and lower runtime.',TEAL)
card(s,9.15,3.45,3.35,1.35,'At 1% support','4,369 frequent itemsets; more combinations survive pruning.',CORAL)
txt(s,9.25,5.25,3.05,.8,'Expected pattern: lowering minimum support increases search space.',11,MID)

# 12 pruning/trim graph
s=newslide('Result 4 — candidate pruning vs transaction trimming','06  /  RESULTS')
add_img(s,GRAPHS/'graph4_candidates_per_pass.png',.55,1.55,9.0,4.85)
card(s,9.65,1.9,2.9,1.45,'Hash pruning','Targets C₂ directly. Collisions limit pruning power.',CORAL)
card(s,9.65,3.65,2.9,1.45,'Transaction trimming','May reduce items scanned later, but adds per-pass bookkeeping.',BLUE)
txt(s,9.7,5.55,2.8,.6,'The trimming rule here is simplified for teaching.',10,MID)

# 13 bucket graph
s=newslide('Result 5 — bucket count changes the trade-off','06  /  RESULTS')
add_img(s,GRAPHS/'graph5_bucket_sweep.png',.55,1.55,9.0,4.85)
card(s,9.65,1.9,2.9,1.45,'Smaller B','More collisions can retain more candidate pairs.',GOLD)
card(s,9.65,3.65,2.9,1.45,'Larger B','Finer buckets may prune more, but use more table memory.',TEAL)
txt(s,9.7,5.55,2.8,.65,'Runtime is not monotonic: hashing and memory overhead matter too.',10,MID)

# 14 conclusions and limits
s=newslide('Takeaways and limitations','07  /  CONCLUSION')
card(s,.8,1.65,3.8,1.8,'What the run shows','On this synthetic workload, Apriori is competitive and often faster; DHP hash-only reduces traced memory in the tested settings.',TEAL)
card(s,4.8,1.65,3.8,1.8,'What DHP contributes','A principled early filter for C₂, with potential gains when pair candidates are the bottleneck.',CORAL)
card(s,8.8,1.65,3.8,1.8,'Why trim did not win','The simplified trim logic and extra processing may cost more than the work saved at this scale.',BLUE)
txt(s,.9,4.05,5.2,.35,'Limitations to disclose',17,NAVY,True)
rich_lines(s,.95,4.55,11.4,1.65,['Synthetic data only; results may not generalize to real transaction distributions.','Only a curated subset completed: 2k/5k/10k size, 2 support thresholds, 3 bucket sizes. The provided 20k file is included, but its full benchmark was not completed.','Single-run timings are exploratory. Full study should use more repeats and verify every bucket-sweep output against Apriori.','Benchmark measures frequent itemsets; confidence-based rule generation is not part of the supplied experiment code.'],12,DARK,7)

# 15 references
s=newslide('Research references','08  /  REFERENCES')
refs=[
('1','Agrawal, R. & Srikant, R. (1994). “Fast Algorithms for Mining Association Rules in Large Databases.” VLDB, 487–499.','https://www.vldb.org/conf/1994/P487.PDF'),
('2','Park, J. S., Chen, M.-S., & Yu, P. S. (1995). “An Effective Hash-Based Algorithm for Mining Association Rules.” SIGMOD ’95, 175–186. DOI: 10.1145/223784.223813.','https://doi.org/10.1145/223784.223813'),
('3','Han, J., Pei, J., & Yin, Y. (2000). “Mining Frequent Patterns without Candidate Generation.” SIGMOD 2000, 1–12. DOI: 10.1145/342009.335372.','https://doi.org/10.1145/342009.335372')]
for i,(num,desc,url) in enumerate(refs):
    y=1.65+i*1.45
    shape(s,.85,y,11.7,1.12,LIGHT,True)
    shape(s,1.05,y+.2,.5,.5,TEAL if i==0 else (CORAL if i==1 else BLUE),True)
    txt(s,1.05,y+.31,.5,.2,num,12,WHITE,True,align=PP_ALIGN.CENTER)
    txt(s,1.8,y+.17,10.35,.53,desc,12,NAVY,True)
    txt(s,1.8,y+.75,10.35,.2,url,9,BLUE)
txt(s,.95,6.35,11.5,.3,'References cover the foundational Apriori and DHP papers plus a widely cited alternative frequent-pattern approach.',10,MID)

prs.save(OUT)
print(f'Saved {OUT}')
